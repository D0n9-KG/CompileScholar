# -*- coding: utf-8 -*-
"""构建方法章 Word 文档：md 源 -> docx（学术样式，三线表，图注题注规范）。
样式：正文宋体小四(12pt) 1.5倍行距首行缩进2字符；章标题黑体16pt；节黑体14pt；
表题/图注黑体小五(9pt)居中；表内宋体10.5pt；三线表无竖线；A4。"""
import io, os, re, sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = f"{HERE}/method_chapter_src.md"
FIG = f"{HERE}/fig_method_arch.png"
OUT = f"{HERE}/METHOD-CHAPTER-draft-v1.docx"


def set_font(run, name_cn, size_pt, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=12, cn="宋体", bold=False, align=None, indent=True,
         space_after=6, line=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    r = p.add_run(text)
    set_font(r, cn, size, bold)
    return p


def heading(doc, text, level):
    sizes = {1: 16, 2: 14, 3: 12.5}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(8 if level == 1 else 6)
    pf.line_spacing = 1.3
    r = p.add_run(text)
    set_font(r, "黑体", sizes[level], bold=(level <= 2))
    # outline level for navigation pane
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level - 1))
    pPr.append(ol)
    return p


def set_cell_border_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def three_line_table(doc, caption, headers, rows, widths=None):
    para(doc, caption, size=9, cn="黑体", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=4)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = 1  # center
    # kill all borders then draw three lines
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)

    def hline(row_idx, edge, sz):
        for cell in t.rows[row_idx].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = tcPr.find(qn("w:tcBorders"))
            if b is None:
                b = OxmlElement("w:tcBorders"); tcPr.append(b)
            old = b.find(qn(f"w:{edge}"))
            if old is not None:
                b.remove(old)
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
            b.append(el)

    allrows = [headers] + rows
    for ri, rowdata in enumerate(allrows):
        # cantSplit
        trPr = t.rows[ri]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for ci, val in enumerate(rowdata):
            cell = t.rows[ri].cells[ci]
            set_cell_border_none(cell)   # kill style-inherited grid per cell
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            set_font(r, "宋体", 10.5, bold=(ri == 0))
            if ci == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hline(0, "top", 12)
    hline(0, "bottom", 6)
    hline(len(rows), "bottom", 12)
    if widths:
        for ci, w in enumerate(widths):
            for ri in range(len(allrows)):
                t.rows[ri].cells[ci].width = Cm(w)
    para(doc, "", size=6, indent=False, space_after=2)


def add_figure(doc, img, caption, width_cm=15.2):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0   # 内嵌图段落必须 1.0 行距
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run()
    r.add_picture(img, width=Cm(width_cm))
    para(doc, caption, size=9, cn="黑体", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=10)


def parse_table_block(inner):
    """【表 3-1：caption。row; row; ...】-> (caption, rows2col)"""
    m = re.match(r"表\s*([\d\-]+)：(.+)", inner, re.S)
    num, rest = m.group(1), m.group(2)
    # caption = up to first '。'
    cap_end = rest.find("。")
    caption_body = rest[:cap_end] if cap_end > 0 else rest[:40]
    rows_src = rest[cap_end + 1:] if cap_end > 0 else ""
    rows = []
    # split only where '；' is followed by a row-head pattern (word + =),
    # so in-row semicolons (e.g. inside config/lineage descriptions) survive
    chunks = re.split(r"；(?=\s*(?:[A-Za-z0-9_一-鿿]{1,16}=|[一二三四五六七八九十]+、))", rows_src)
    for chunk in chunks:
        chunk = chunk.strip().rstrip("。").strip()
        if not chunk:
            continue
        # strip leading ordinals 一、二、...
        chunk = re.sub(r"^[一二三四五六七八九十]+、\s*", "", chunk)
        m = re.match(r"^([^=：]{1,14})：(.+)$", chunk, re.S)
        if m:  # name：description form (表3-3) — colon wins over inner '='
            rows.append([m.group(1).strip(), m.group(2).strip()])
        elif "=" in chunk:
            k, v = chunk.split("=", 1)
            rows.append([k.strip(), v.strip()])
        else:
            rows.append([chunk, ""])
    return f"表 {num}　{caption_body}", rows


def main():
    text = io.open(SRC, encoding="utf-8").read()
    doc = Document()
    # page setup A4
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)

    lines = text.split("\n")
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        if s.startswith("### "):
            heading(doc, s[4:], 3)
        elif s.startswith("## "):
            heading(doc, s[3:], 2)
        elif s.startswith("# "):
            heading(doc, s[2:], 1)
        elif s.startswith("【图"):
            inner = s.strip("【】")
            m = re.match(r"图\s*([\d\-]+)：(.+)", inner, re.S)
            cap = f"图 {m.group(1)}　{m.group(2)}"
            add_figure(doc, FIG, cap)
        elif s.startswith("【表"):
            inner = s.strip("【】")
            caption, rows = parse_table_block(inner)
            hdr = rows[0] if rows and rows[0][1] == "" and False else None
            if caption.startswith("表 3-1"):
                headers = ["记录类型", "语义与必填字段"]
            elif caption.startswith("表 3-2"):
                headers = ["工具", "语义与对应访问负载"]
            else:
                headers = ["招", "动作规范与所消费的表示语义"]
            three_line_table(doc, caption, headers, rows, widths=[3.2, 12.0])
        else:
            para(doc, s)

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
